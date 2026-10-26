import openpyxl, re, json, sys, datetime as dt, unicodedata, os, glob
U = sys.argv[1]; OUT = sys.argv[2]

def norm(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii','ignore').decode().lower()
    return re.sub(r'\s+', ' ', s).strip()

def num(v):
    if v is None or isinstance(v, bool): return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace(',', '')
    try: return float(s)
    except: return None

def txt(v):
    if v is None: return ''
    if isinstance(v, (int, float)): return ('%g' % v)
    return str(v).strip()

# ---------------- CONTROL CIBER ----------------
CAMPOS = {  # encabezado normalizado -> campo
 'dia':'venta', 'sencillo':'sencillo', 'saldo':'saldo', 'recargas':'recargas', 'caja':'caja',
 'ingresos':'ingresos', 'gastos':'gastos', 'pagos':'pagos', 'gasto banca':'gasto_banca',
 'banca':'banca', 'billete total':'banca', 'sencillo total':'sencillo_total', 'cambio':'cambio',
 'sencillo a billete':'sen_a_bil', 'billete a sencillo':'bil_a_sen', 'otros ingresos':'otros_ingresos',
 'total pdv':'total_pdv', 'total global':'total_global',
}
CON_TEXTO = {'ingresos','gastos','pagos','gasto_banca','otros_ingresos'}

def control_ciber(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    dias = {}
    for ws in wb.worksheets:
        if not re.fullmatch(r'\d{4}', ws.title): continue
        anio = int(ws.title)
        rows = ws.iter_rows(min_row=1, max_row=400, max_col=40, values_only=True)
        head = next(rows)
        col = {}; billetes = None
        for i, h in enumerate(head):
            if h is None: continue
            n = norm(h)
            if n.startswith('billetes hoy'): billetes = i; continue
            if n in CAMPOS and CAMPOS[n] not in col: col[CAMPOS[n]] = i
        prev = None
        for r in rows:
            if r is None or len(r) == 0: continue
            f = r[0]
            if isinstance(f, dt.datetime) and f.year == anio: d = f.date()
            elif prev is not None and (f is not None or any(x is not None for x in r[1:12])): d = prev + dt.timedelta(days=1)
            elif isinstance(f, dt.datetime) and prev is None: d = dt.date(anio, 1, 1)
            else: continue
            if prev is not None and d <= prev: d = prev + dt.timedelta(days=1)
            if d.year != anio: break
            prev = d
            x = {'fecha': d.isoformat()}
            for k, i in col.items():
                v = r[i] if i < len(r) else None
                x[k] = num(v)
                if k in CON_TEXTO and i + 1 < len(r):
                    t = r[i + 1]
                    if isinstance(t, str) and t.strip(): x[k + '_txt'] = t.strip()
            if billetes is not None:
                x['billetes'] = num(r[billetes])
                for j, den in enumerate([20, 50, 100, 200, 500, 1000]):
                    v = num(r[billetes + 1 + j]) if billetes + 1 + j < len(r) else None
                    if v: x['b%d' % den] = int(v)
            if isinstance(r[1], str) and r[1].strip() and num(r[1]) is None:
                x['notas'] = r[1].strip()   # p. ej. "EDAD OSCURA"
            # obs de gastos/pagos en hojas 2010-2012 (columna "Observaciones" tras gastos)
            util = [k for k, v in x.items() if k != 'fecha' and v not in (None, 0, '')]
            if not util: continue
            dias[x['fecha']] = {k: v for k, v in x.items() if v not in (None, '')}
    return dias

# ---------------- RECARGAS ----------------
RX_FECHA = [
  (re.compile(r'^(\d{1,2})/(\d{1,2})/(\d{4})\s+(\d{1,2}):(\d{2})(?::(\d{2}))?\s*([ap])?\.?\s*m?', re.I), 'dmy'),
  (re.compile(r'^(\d{1,2})\s+([a-z]{3})\w*\.?\s+(\d{4}),?\s+(\d{1,2}):(\d{2})', re.I), 'dMy'),
]
MESES = dict(ene=1,feb=2,mar=3,abr=4,may=5,jun=6,jul=7,ago=8,sep=9,oct=10,nov=11,dic=12,
             jan=1,apr=4,aug=8,dec=12)
def fecha(v):
    if isinstance(v, dt.datetime): return v
    if not isinstance(v, str): return None
    s = v.strip().lower()
    m = RX_FECHA[0][0].match(s)
    if m:
        d, mo, y, h, mi, se, ap = m.groups()
        h = int(h)
        if ap == 'p' and h < 12: h += 12
        if ap == 'a' and h == 12: h = 0
        try: return dt.datetime(int(y), int(mo), int(d), h, int(mi), int(se or 0))
        except: return None
    m = RX_FECHA[1][0].match(s)
    if m:
        d, mo, y, h, mi = m.groups()
        if mo[:3] not in MESES: return None
        return dt.datetime(int(y), MESES[mo[:3]], int(d), int(h), int(mi))
    return None

NO_CIA = {'a pdf','exito','web','app','android','ciber dark','ver ticket','ios','transaccion realizada','00 respuesta ok'}
def es_tel(v):
    if isinstance(v, float) and v.is_integer(): v = int(v)
    s = str(v).strip() if v is not None else ''
    return s if re.fullmatch(r'\d{10}', s) else None

def recargas(path, proveedor, out, abonos, nombres):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        if ws.title.upper() == 'NOMBRES':
            for r in ws.iter_rows(min_row=2, max_col=2, values_only=True):
                t = es_tel(r[0]) if r and r[0] is not None else None
                if t and r[1] and str(r[1]).strip() and str(r[1]).strip() != '#N/A':
                    nombres.setdefault(t, str(r[1]).strip())
            continue
        for r in ws.iter_rows(min_row=2, max_col=46, values_only=True):
            if not r: continue
            # abono de saldo del día: columna H(7)/I(8) = fecha del día, M(12) = abono
            for ci, cm in ((7, 12), (8, 12)):
                if len(r) > cm and isinstance(r[ci], dt.datetime) and num(r[cm]) and num(r[cm]) > 0:
                    abonos.append({'fecha': r[ci].date().isoformat(), 'monto': num(r[cm]), 'proveedor': proveedor})
            cel = list(r[12:])
            p = next((i for i, v in enumerate(cel) if es_tel(v)), None)
            d = next((i for i, v in enumerate(cel) if fecha(v)), None)
            if p is None or d is None: continue
            f = fecha(cel[d]); tel = es_tel(cel[p])
            viejo = any(isinstance(v, str) and norm(v) in ('a pdf', 'exito', 'ciber dark') for v in cel)
            monto = num(cel[p + 1]) if viejo and p + 1 < len(cel) else None
            if monto is None: monto = num(cel[p - 1]) if p > 0 else None
            if monto is None or monto <= 0: continue
            if viejo:
                cia = txt(cel[d + 1]) if d + 1 < len(cel) else ''
                folio = txt(cel[p + 2]) if p + 2 < len(cel) else ''
                estado = txt(cel[p + 3]) if p + 3 < len(cel) else ''
                desc = next((txt(v) for v in cel[p + 4:] if isinstance(v, str) and re.search(r'[a-z].*\d', v, re.I)
                             and norm(v) not in NO_CIA and not norm(v).startswith('transacc') and not norm(v).startswith('00 ')), '')
                idp = txt(cel[d - 1]) if d > 0 else ''
            else:
                cia = next((txt(v) for v in cel if isinstance(v, str) and re.fullmatch(r'[A-Za-z_ ]{2,30}', v.strip())
                            and norm(v) not in NO_CIA), '')
                ci = next((i for i, v in enumerate(cel) if isinstance(v, str) and txt(v) == cia), None)
                cand = [cel[i] for i in (p + 1, (ci or -9) + 1) if 0 <= i < len(cel) and i != d]
                folio = next((txt(v) for v in cand if v is not None and fecha(v) is None and norm(v) not in NO_CIA), '')
                estado = ''; desc = ''; idp = ''
            if estado and norm(estado) not in ('exito', ''): continue
            out.append({'fecha': f.isoformat(sep=' ', timespec='minutes'), 'compania': cia.replace('_', ' ').strip().upper(),
                        'telefono': tel, 'monto': monto, 'folio': folio, 'producto': desc, 'proveedor': proveedor, 'id_prov': idp})

def main():
    archivos = sorted(glob.glob(os.path.join(U, '*.xls*')))
    cc = next(a for a in archivos if 'Control_Ciber' in a)
    dias = control_ciber(cc)
    rec, abonos, nombres = [], [], {}
    for a in archivos:
        b = os.path.basename(a)
        if 'TAEPOCKET' in b: recargas(a, 'TAEPOCKET', rec, abonos, nombres)
        elif 'RECARGAS' in b: recargas(a, 'RECARGAS MAS', rec, abonos, nombres)
    vistos = set(); rec2 = []
    for x in sorted(rec, key=lambda x: x['fecha']):
        k = (x['proveedor'], x['fecha'], x['telefono'], x['monto'], x['folio'])
        if k in vistos: continue
        vistos.add(k); rec2.append(x)
    ab = {}
    for x in abonos: ab[(x['fecha'], x['proveedor'], x['monto'])] = x
    data = {'version': 1, 'dias': [dias[k] for k in sorted(dias)], 'recargas': rec2,
            'abonos': sorted(ab.values(), key=lambda x: x['fecha']),
            'nombres': [{'telefono': t, 'nombre': n} for t, n in sorted(nombres.items())]}
    json.dump(data, open(OUT, 'w'), ensure_ascii=False, separators=(',', ':'))
    print('dias', len(data['dias']), 'recargas', len(rec2), 'abonos', len(data['abonos']), 'nombres', len(nombres))
main()
